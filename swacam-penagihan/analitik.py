"""
Analitik & laporan penagihan tunggakan.

  GET /api/analitik                          -> tren per bulan + efektivitas template (JSON, untuk dashboard)
  GET /api/analitik/laporan/excel?periode=   -> laporan .xlsx (ringkasan, perbandingan antarbulan, template, detail)
  GET /api/analitik/laporan/pdf?periode=     -> laporan .pdf untuk atasan

Sumber data:
  - tunggakan_pelanggan : status bayar tiap pelanggan per periode (bulan). Tiap upload file memperbarui statusnya.
  - wa_riwayat          : setiap pesan WA (terkirim / diterima / dibaca / balasan) beserta template yang dipakai.

Catatan penting (juga ditampilkan di laporan):
  - "% lunas" = pelanggan berstatus LUNAS pada file terakhir yang memuat periode itu / seluruh pelanggan di periode itu.
  - "Konversi template" = pelanggan yang lunas SETELAH menerima template tersebut (template terakhir yang diterima sebelum
    lunas). Ini hubungan waktu, bukan bukti sebab-akibat: pelanggan bisa saja membayar tanpa dipengaruhi pesan.
"""
import datetime
import io
import re
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from sqlalchemy import text

from auth import get_current_admin, get_current_admin_from_query
from db_handler import SessionLocal, TunggakanPelanggan, log_activity
from tunggakan import BULAN_ID, rupiah

router = APIRouter(prefix="/api/analitik", tags=["Analitik & Laporan"])

BASE_DIR = Path(__file__).resolve().parent
MIN_DATA_TEMPLATE = 5          # template dengan pelanggan < 5 ditandai "data masih sedikit"
CATATAN = [
    "% Lunas = pelanggan berstatus lunas pada file terakhir yang diupload untuk bulan tersebut, dibagi seluruh pelanggan bulan itu.",
    "Tingkat baca dan balasan WhatsApp hanya terisi bila webhook Fonnte aktif; tanpa webhook status berhenti di 'Terkirim'.",
    "Konversi template = pelanggan yang lunas setelah menerima template tersebut (template terakhir sebelum lunas). "
    "Ini hubungan waktu, bukan bukti bahwa pesan menjadi penyebab pembayaran.",
]


# ─────────────────────────────────────────────────────────────
#  Helper
# ─────────────────────────────────────────────────────────────
def urut_periode(label: Optional[str]) -> int:
    """'Agustus 2026' -> 202608 (0 bila tidak dikenali) untuk pengurutan kronologis."""
    m = re.fullmatch(r"\s*([A-Za-z]+)\s+(\d{4})\s*", label or "")
    if not m:
        return 0
    nama = m.group(1).capitalize()
    return int(m.group(2)) * 100 + BULAN_ID.index(nama) + 1 if nama in BULAN_ID else 0


def _pct(a, b):
    return round(a / b * 100, 1) if b else None


def _selisih(a, b):
    return round(a - b, 1) if a is not None and b is not None else None


def _kategori(t: dict, today: datetime.date) -> str:
    if (t.get("status_bayar") or "BELUM") == "LUNAS":
        return "LUNAS"
    jt = t.get("jatuh_tempo")
    if jt and today <= jt:
        return "BELUM_JATUH_TEMPO"
    return "MENUNGGAK"


# ─────────────────────────────────────────────────────────────
#  Perhitungan (fungsi murni: mudah diuji tanpa database)
# ─────────────────────────────────────────────────────────────
def hitung_ringkasan(t_rows: List[dict], w_rows: List[dict], today: datetime.date = None) -> List[dict]:
    """Satu baris per periode (urut kronologis): data pembayaran + statistik WhatsApp + selisih vs bulan sebelumnya."""
    today = today or datetime.date.today()
    per: Dict[str, dict] = {}

    def slot(p):
        p = p or "Tanpa periode"
        return per.setdefault(p, dict(
            periode=p, urut=urut_periode(p), total=0, lunas=0, menunggak=0, belum_jatuh_tempo=0,
            nominal_total=0.0, nominal_lunas=0.0, nominal_belum_lunas=0.0,
            wa_terkirim=0, wa_diterima=0, wa_dibaca=0, wa_balasan=0, wa_gagal=0, pelanggan_dihubungi=set()))

    for t in t_rows:
        s = slot(t.get("periode"))
        nom = float(t.get("nominal") or 0)
        kat = _kategori(t, today)
        s["total"] += 1
        s["nominal_total"] += nom
        if kat == "LUNAS":
            s["lunas"] += 1
            s["nominal_lunas"] += nom
        else:
            s["nominal_belum_lunas"] += nom
            s["menunggak" if kat == "MENUNGGAK" else "belum_jatuh_tempo"] += 1

    for w in w_rows:
        s = slot(w.get("periode"))
        st = (w.get("status") or "").upper()
        if st == "GAGAL":
            s["wa_gagal"] += 1
            continue
        if st == "ANTRE":
            continue
        s["wa_terkirim"] += 1
        s["pelanggan_dihubungi"].add(w.get("id_pelanggan"))
        if st in ("DITERIMA", "DIBACA") or w.get("waktu_diterima") or w.get("waktu_dibaca"):
            s["wa_diterima"] += 1
        if st == "DIBACA" or w.get("waktu_dibaca"):
            s["wa_dibaca"] += 1
        if w.get("balasan"):
            s["wa_balasan"] += 1

    hasil = sorted(per.values(), key=lambda x: (x["urut"], x["periode"]))
    for s in hasil:
        s["pelanggan_dihubungi"] = len(s["pelanggan_dihubungi"])
        s["persen_lunas"] = _pct(s["lunas"], s["total"])
        s["persen_lunas_nominal"] = _pct(s["nominal_lunas"], s["nominal_total"])
        s["persen_diterima"] = _pct(s["wa_diterima"], s["wa_terkirim"])
        s["persen_baca"] = _pct(s["wa_dibaca"], s["wa_terkirim"])
        s["persen_balasan"] = _pct(s["wa_balasan"], s["wa_terkirim"])
    for i, s in enumerate(hasil):
        p = hasil[i - 1] if i else None
        s["delta"] = {k: (_selisih(s[k], p[k]) if p else None)
                      for k in ("persen_lunas", "persen_baca", "persen_balasan")}
        s["delta"]["nominal_belum_lunas"] = round(s["nominal_belum_lunas"] - p["nominal_belum_lunas"]) if p else None
    return hasil


def hitung_template(t_rows: List[dict], w_rows: List[dict], info_tpl: Dict[str, dict] = None) -> List[dict]:
    """Efektivitas tiap template: dibaca, dibalas, dan lunas setelah menerima template itu."""
    info_tpl = info_tpl or {}
    lunas = {(t["id_pelanggan"], t.get("periode") or ""): t
             for t in t_rows if (t.get("status_bayar") or "BELUM") == "LUNAS"}
    pesan_per_pelanggan = defaultdict(list)
    for w in w_rows:
        if (w.get("status") or "").upper() in ("GAGAL", "ANTRE"):
            continue
        pesan_per_pelanggan[(w["id_pelanggan"], w.get("periode") or "")].append(w)

    st: Dict[str, dict] = {}

    def slot(kode):
        return st.setdefault(kode, dict(kode=kode, pesan=0, dibaca=0, balasan=0, pelanggan=set(), lunas=set(), hari=[]))

    for key, msgs in pesan_per_pelanggan.items():
        msgs.sort(key=lambda m: m.get("waktu_kirim") or datetime.datetime.min)
        for m in msgs:
            s = slot(m.get("template_kode") or "standar")
            s["pesan"] += 1
            s["pelanggan"].add(key)
            if (m.get("status") or "").upper() == "DIBACA" or m.get("waktu_dibaca"):
                s["dibaca"] += 1
            if m.get("balasan"):
                s["balasan"] += 1
        t = lunas.get(key)
        if t and t.get("lunas_pada"):
            sebelum = [m for m in msgs if m.get("waktu_kirim") and m["waktu_kirim"] <= t["lunas_pada"]]
            if sebelum:
                terakhir = sebelum[-1]
                s = slot(terakhir.get("template_kode") or "standar")
                s["lunas"].add(key)
                s["hari"].append((t["lunas_pada"] - terakhir["waktu_kirim"]).total_seconds() / 86400)

    hasil = []
    for kode, s in st.items():
        info = info_tpl.get(kode) or {}
        n_pel = len(s["pelanggan"])
        hasil.append(dict(
            kode=kode, nama=info.get("nama") or ("Pesan standar" if kode == "standar" else kode),
            tahap=info.get("tahap"), pesan=s["pesan"], pelanggan=n_pel,
            dibaca=s["dibaca"], balasan=s["balasan"], lunas=len(s["lunas"]),
            persen_baca=_pct(s["dibaca"], s["pesan"]), persen_balasan=_pct(s["balasan"], s["pesan"]),
            konversi=_pct(len(s["lunas"]), n_pel),
            rata_hari=round(sum(s["hari"]) / len(s["hari"]), 1) if s["hari"] else None,
            data_cukup=n_pel >= MIN_DATA_TEMPLATE, terbaik=False))
    hasil.sort(key=lambda x: (-(x["konversi"] or 0), -x["pesan"]))
    layak = [h for h in hasil if h["data_cukup"] and (h["konversi"] or 0) > 0]
    if layak:
        layak[0]["terbaik"] = True
    return hasil


def baris_detail(t_rows: List[dict], periode: str, today: datetime.date = None) -> List[dict]:
    """Daftar pelanggan pada satu periode (menunggak paling lama / terbesar di atas)."""
    today = today or datetime.date.today()
    out = []
    for t in t_rows:
        if (t.get("periode") or "Tanpa periode") != periode:
            continue
        kat = _kategori(t, today)
        jt = t.get("jatuh_tempo")
        out.append(dict(t, kategori=kat, hari_terlambat=(today - jt).days if kat == "MENUNGGAK" and jt else None))
    out.sort(key=lambda x: ({"MENUNGGAK": 0, "BELUM_JATUH_TEMPO": 1, "LUNAS": 2}[x["kategori"]],
                            -(x["hari_terlambat"] or 0), -(x.get("nominal") or 0)))
    return out


# ─────────────────────────────────────────────────────────────
#  Ambil data dari database
# ─────────────────────────────────────────────────────────────
def _dt(v):
    """Pastikan nilai waktu berupa datetime (MySQL sudah datetime; driver lain bisa mengembalikan teks)."""
    if v is None or isinstance(v, datetime.datetime):
        return v
    try:
        return datetime.datetime.fromisoformat(str(v).replace("Z", ""))
    except ValueError:
        return None


def muat_data():
    s = SessionLocal()
    try:
        t = [dict(id_pelanggan=x.id_pelanggan, nama=x.nama, no_hp=x.no_hp, periode=x.periode, nominal=x.nominal,
                  status_bayar=x.status_bayar, jatuh_tempo=x.jatuh_tempo, lunas_pada=x.lunas_pada,
                  wa_status=x.wa_status, wa_jumlah_kirim=x.wa_jumlah_kirim)
             for x in s.query(TunggakanPelanggan).all()]
        w, info = [], {}
        try:
            w = [dict(r) for r in s.execute(text(
                "SELECT id_pelanggan, periode, template_kode, status, waktu_kirim, waktu_diterima, waktu_dibaca, balasan "
                "FROM wa_riwayat")).mappings().all()]
            for r in w:
                for k in ("waktu_kirim", "waktu_diterima", "waktu_dibaca"):
                    r[k] = _dt(r[k])
        except Exception:
            s.rollback()
        try:
            info = {r["kode"]: dict(r) for r in s.execute(text("SELECT kode, nama, tahap FROM wa_template")).mappings().all()}
        except Exception:
            s.rollback()
        return t, w, info
    finally:
        s.close()


def susun_analitik(t, w, info, today=None) -> dict:
    bulan = hitung_ringkasan(t, w, today)
    return {
        "bulan": bulan,
        "template": hitung_template(t, w, info),
        "webhook_aktif": any(b["wa_diterima"] or b["wa_dibaca"] or b["wa_balasan"] for b in bulan),
        "catatan": CATATAN,
        "dibuat": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
    }


@router.get("")
def analitik(admin: dict = Depends(get_current_admin)):
    t, w, info = muat_data()
    return susun_analitik(t, w, info)


# ─────────────────────────────────────────────────────────────
#  Laporan Excel
# ─────────────────────────────────────────────────────────────
def _pilih_periode(bulan: List[dict], periode: Optional[str]) -> Optional[dict]:
    if not bulan:
        return None
    if periode:
        for b in bulan:
            if b["periode"] == periode:
                return b
        raise HTTPException(404, f"Periode '{periode}' tidak ditemukan.")
    return bulan[-1]


def bangun_excel(data: dict, detail: List[dict], periode: Optional[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    F = "Arial"
    biru = PatternFill("solid", fgColor="003C88")
    abu = PatternFill("solid", fgColor="F1F5F9")
    tipis = Side(style="thin", color="CBD5E1")
    garis = Border(left=tipis, right=tipis, top=tipis, bottom=tipis)
    hd_font = Font(name=F, bold=True, color="FFFFFF", size=10)
    wb = Workbook()

    def header(ws, row, judul):
        for i, j in enumerate(judul, 1):
            c = ws.cell(row=row, column=i, value=j)
            c.font, c.fill, c.border = hd_font, biru, garis
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 32

    def lebar(ws, ukuran):
        for i, w in enumerate(ukuran, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    bulan = data["bulan"]
    # ── Sheet 2: Perbandingan Bulan (dibuat dulu karena Ringkasan merujuk ke sini) ──
    wp = wb.active
    wp.title = "Perbandingan Bulan"
    wp["A1"] = "Perbandingan Antarbulan"
    wp["A1"].font = Font(name=F, bold=True, size=14, color="003C88")
    wp["A2"] = "Nilai persen dihitung dengan rumus dari kolom jumlah di sebelah kirinya."
    wp["A2"].font = Font(name=F, italic=True, size=9, color="64748B")
    judul = ["Periode", "Total Pelanggan", "Lunas", "Menunggak", "Belum Jatuh Tempo", "Nominal Total (Rp)",
             "Nominal Lunas (Rp)", "Nominal Belum Lunas (Rp)", "% Lunas", "% Lunas (Nominal)", "WA Terkirim",
             "WA Dibaca", "WA Balasan", "% Dibaca", "% Balasan", "Perubahan % Lunas"]
    header(wp, 4, judul)
    r0 = 5
    for i, b in enumerate(bulan):
        r = r0 + i
        vals = [b["periode"], b["total"], b["lunas"], b["menunggak"], b["belum_jatuh_tempo"], round(b["nominal_total"]),
                round(b["nominal_lunas"]), f"=F{r}-G{r}", f"=IF(B{r}=0,0,C{r}/B{r})", f"=IF(F{r}=0,0,G{r}/F{r})",
                b["wa_terkirim"], b["wa_dibaca"], b["wa_balasan"], f"=IF(K{r}=0,0,L{r}/K{r})", f"=IF(K{r}=0,0,M{r}/K{r})",
                f"=I{r}-I{r - 1}" if i else "-"]
        for j, v in enumerate(vals, 1):
            c = wp.cell(row=r, column=j, value=v)
            c.font, c.border = Font(name=F, size=10), garis
            if j in (6, 7, 8):
                c.number_format = '#,##0'
            elif j in (9, 10, 14, 15):
                c.number_format = '0.0%'
            elif j == 16:
                c.number_format = '+0.0%;-0.0%;0.0%'
                c.alignment = Alignment(horizontal="right")
    rn = r0 + len(bulan) - 1
    rt = rn + 1
    if bulan:
        tot = ["TOTAL"] + [f"=SUM({get_column_letter(j)}{r0}:{get_column_letter(j)}{rn})" for j in (2, 3, 4, 5, 6, 7, 8)]
        tot += [f"=IF(B{rt}=0,0,C{rt}/B{rt})", f"=IF(F{rt}=0,0,G{rt}/F{rt})"]
        tot += [f"=SUM({get_column_letter(j)}{r0}:{get_column_letter(j)}{rn})" for j in (11, 12, 13)]
        tot += [f"=IF(K{rt}=0,0,L{rt}/K{rt})", f"=IF(K{rt}=0,0,M{rt}/K{rt})", ""]
        for j, v in enumerate(tot, 1):
            c = wp.cell(row=rt, column=j, value=v)
            c.font, c.fill, c.border = Font(name=F, bold=True, size=10), abu, garis
            c.number_format = '#,##0' if j in (6, 7, 8) else ('0.0%' if j in (9, 10, 14, 15) else 'General')
        lebar(wp, [18, 12, 10, 12, 14, 18, 18, 20, 10, 12, 11, 11, 11, 11, 11, 14])
        wp.freeze_panes = "B5"
        ch = LineChart()
        ch.title, ch.height, ch.width = "% Lunas per Bulan", 8, 16
        ch.add_data(Reference(wp, min_col=9, min_row=4, max_row=rn), titles_from_data=True)
        ch.set_categories(Reference(wp, min_col=1, min_row=r0, max_row=rn))
        ch.y_axis.number_format = '0%'
        ch.y_axis.delete = ch.x_axis.delete = False
        wp.add_chart(ch, f"A{rt + 3}")
        bc = BarChart()
        bc.title, bc.height, bc.width = "Nominal Belum Lunas per Bulan (Rp)", 8, 16
        bc.add_data(Reference(wp, min_col=8, min_row=4, max_row=rn), titles_from_data=True)
        bc.set_categories(Reference(wp, min_col=1, min_row=r0, max_row=rn))
        bc.y_axis.delete = bc.x_axis.delete = False
        wp.add_chart(bc, f"H{rt + 3}")
    else:
        wp["A5"] = "Belum ada data."

    # ── Sheet 1: Ringkasan ──
    ws = wb.create_sheet("Ringkasan", 0)
    ws["A1"] = "LAPORAN TUNGGAKAN & EFEKTIVITAS PENAGIHAN WHATSAPP"
    ws["A1"].font = Font(name=F, bold=True, size=15, color="003C88")
    ws["A2"] = f"Periode: {periode['periode'] if periode else '-'}   |   Dibuat: {data['dibuat']}"
    ws["A2"].font = Font(name=F, size=10, color="475569")
    if periode:
        idx = next(i for i, b in enumerate(bulan) if b["periode"] == periode["periode"])
        r = r0 + idx
        kp = "'Perbandingan Bulan'!"
        baris = [
            ("Total pelanggan", f"={kp}B{r}", "#,##0"), ("Sudah lunas", f"={kp}C{r}", "#,##0"),
            ("Menunggak (lewat jatuh tempo)", f"={kp}D{r}", "#,##0"), ("Belum jatuh tempo", f"={kp}E{r}", "#,##0"),
            ("Nominal belum lunas (Rp)", f"={kp}H{r}", "#,##0"), ("% Lunas", f"={kp}I{r}", "0.0%"),
            ("% Lunas (nominal)", f"={kp}J{r}", "0.0%"), ("Pesan WA terkirim", f"={kp}K{r}", "#,##0"),
            ("% Dibaca", f"={kp}N{r}", "0.0%"), ("% Dibalas", f"={kp}O{r}", "0.0%"),
        ]
        header(ws, 4, ["Indikator", "Nilai"])
        for i, (k, f, nf) in enumerate(baris, 5):
            a, b = ws.cell(row=i, column=1, value=k), ws.cell(row=i, column=2, value=f)
            a.font, b.font = Font(name=F, size=10), Font(name=F, size=10, bold=True)
            a.border = b.border = garis
            b.number_format = nf
            b.alignment = Alignment(horizontal="right")
        nr = 5 + len(baris) + 1
    else:
        nr = 4
    ws.cell(row=nr, column=1, value="Catatan & asumsi").font = Font(name=F, bold=True, size=10)
    for i, c in enumerate(data["catatan"], nr + 1):
        x = ws.cell(row=i, column=1, value=f"{i - nr}. {c}")
        x.font = Font(name=F, size=9, color="475569")
        x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
        ws.row_dimensions[i].height = 30
    lebar(ws, [34, 20, 14, 14, 14, 14])

    # ── Sheet 3: Efektivitas Template ──
    wt = wb.create_sheet("Efektivitas Template")
    wt["A1"] = "Efektivitas Template WhatsApp"
    wt["A1"].font = Font(name=F, bold=True, size=14, color="003C88")
    wt["A2"] = "Konversi = pelanggan yang lunas setelah menerima template ini. Hubungan waktu, bukan bukti sebab-akibat."
    wt["A2"].font = Font(name=F, italic=True, size=9, color="64748B")
    header(wt, 4, ["Template", "Tahap", "Pesan Terkirim", "Pelanggan", "Dibaca", "Dibalas", "Lunas Setelahnya",
                   "% Dibaca", "% Dibalas", "% Konversi", "Rata-rata Hari ke Lunas", "Keterangan"])
    for i, t in enumerate(data["template"], 5):
        ket = ("Terbaik" if t["terbaik"] else "") + (" Data masih sedikit" if not t["data_cukup"] else "")
        vals = [t["nama"], t["tahap"] if t["tahap"] is not None else "-", t["pesan"], t["pelanggan"], t["dibaca"],
                t["balasan"], t["lunas"], f"=IF(C{i}=0,0,E{i}/C{i})", f"=IF(C{i}=0,0,F{i}/C{i})",
                f"=IF(D{i}=0,0,G{i}/D{i})", t["rata_hari"] if t["rata_hari"] is not None else "-", ket.strip()]
        for j, v in enumerate(vals, 1):
            c = wt.cell(row=i, column=j, value=v)
            c.font, c.border = Font(name=F, size=10), garis
            if j in (8, 9, 10):
                c.number_format = '0.0%'
    if not data["template"]:
        wt["A5"] = "Belum ada pesan WhatsApp yang terkirim."
    lebar(wt, [38, 8, 14, 12, 10, 10, 16, 10, 10, 11, 18, 26])
    wt.freeze_panes = "B5"

    # ── Sheet 4: Data pelanggan periode terpilih ──
    if periode:
        wd = wb.create_sheet(f"Data {periode['periode']}"[:31])
        wd["A1"] = f"Data Pelanggan - {periode['periode']}"
        wd["A1"].font = Font(name=F, bold=True, size=14, color="003C88")
        header(wd, 3, ["No", "ID Pelanggan", "Nama", "No. WA", "Tagihan (Rp)", "Jatuh Tempo", "Terlambat (hari)",
                       "Pembayaran", "Status WA", "Dikirim (x)"])
        for i, d in enumerate(detail, 1):
            vals = [i, d["id_pelanggan"], d.get("nama") or "-", d.get("no_hp") or "-", d.get("nominal") or 0,
                    d["jatuh_tempo"].strftime("%d-%m-%Y") if d.get("jatuh_tempo") else "-",
                    d["hari_terlambat"] if d["hari_terlambat"] is not None else "-",
                    {"MENUNGGAK": "Menunggak", "BELUM_JATUH_TEMPO": "Belum jatuh tempo", "LUNAS": "Lunas"}[d["kategori"]],
                    d.get("wa_status") or "BELUM", d.get("wa_jumlah_kirim") or 0]
            for j, v in enumerate(vals, 1):
                c = wd.cell(row=3 + i, column=j, value=v)
                c.font, c.border = Font(name=F, size=10), garis
                if j == 5:
                    c.number_format = '#,##0'
        if detail:
            wd.auto_filter.ref = f"A3:J{3 + len(detail)}"
        wd.freeze_panes = "A4"
        lebar(wd, [6, 18, 28, 16, 16, 14, 16, 18, 14, 12])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────
#  Laporan PDF (fpdf2, gaya sama dengan laporan rekap SWACAM)
# ─────────────────────────────────────────────────────────────
def _c(t) -> str:
    return str("" if t is None else t).encode("latin-1", "replace").decode("latin-1")


def _pc(v) -> str:
    return "-" if v is None else f"{v:.1f}%".replace(".", ",")


def bangun_pdf(data: dict, detail: List[dict], periode: Optional[dict]) -> bytes:
    from fpdf import FPDF

    logo = BASE_DIR / "static" / "logo-pln.png"
    judul_periode = periode["periode"] if periode else "Semua periode"
    dibuat = data["dibuat"]
    BIRU, ABU = (0, 60, 136), (241, 245, 249)

    class Lap(FPDF):
        def header(self):
            with self.local_context(font_family="Arial", font_style="B", font_size=60, text_color=(240, 240, 240)):
                with self.rotation(30, self.w / 2, self.h / 2):
                    self.set_xy(0, self.h / 2 - 15)
                    self.cell(self.w, 15, "PLN SWACAM", align="C")
            x = 10
            if logo.exists():
                self.image(str(logo), x=10, y=6, h=12)
                x = 26
            self.set_xy(x, 6)
            self.set_font("Arial", "B", 13)
            self.set_text_color(*BIRU)
            self.cell(0, 6, "LAPORAN TUNGGAKAN & EFEKTIVITAS PENAGIHAN WHATSAPP", new_x="LMARGIN", new_y="NEXT", align="C")
            self.set_font("Arial", "", 8)
            self.set_text_color(90, 90, 90)
            self.cell(0, 4, _c(f"Periode: {judul_periode}   |   Dibuat: {dibuat}"), new_x="LMARGIN", new_y="NEXT", align="C")
            self.set_draw_color(*BIRU)
            self.set_line_width(0.6)
            self.line(10, self.get_y() + 1.5, self.w - 10, self.get_y() + 1.5)
            self.set_y(self.get_y() + 5)

        def footer(self):
            self.set_y(-11)
            self.set_fill_color(*BIRU)
            self.rect(10, self.get_y(), self.w - 20, 7, "F")
            self.set_font("Arial", "B", 6.5)
            self.set_text_color(255, 255, 255)
            self.set_xy(10, self.get_y() + 1.5)
            self.cell((self.w - 20) / 2, 4, "PLN - Laporan Tunggakan & Penagihan WA", align="L")
            self.cell((self.w - 20) / 2, 4, f"Hal. {self.page_no()}", align="R")

    pdf = Lap("L", "mm", "A4")
    pdf.set_margins(10, 28, 10)
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()

    def judul(t):
        pdf.set_font("Arial", "B", 10)
        pdf.set_text_color(*BIRU)
        pdf.cell(0, 7, _c(t), new_x="LMARGIN", new_y="NEXT")

    def tabel(kolom, baris, tinggi=6.5, ukuran=7):
        pdf.set_fill_color(*BIRU)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Arial", "B", ukuran)
        for n, w, _ in kolom:
            pdf.cell(w, 7.5, _c(n), border=1, align="C", fill=True)
        pdf.ln()
        pdf.set_text_color(40, 40, 40)
        for i, r in enumerate(baris):
            if pdf.get_y() > pdf.h - 24:
                pdf.add_page()
                pdf.set_fill_color(*BIRU)
                pdf.set_text_color(255, 255, 255)
                pdf.set_font("Arial", "B", ukuran)
                for n, w, _ in kolom:
                    pdf.cell(w, 7.5, _c(n), border=1, align="C", fill=True)
                pdf.ln()
                pdf.set_text_color(40, 40, 40)
            pdf.set_fill_color(*(ABU if i % 2 else (255, 255, 255)))
            pdf.set_font("Arial", "", ukuran)
            for (n, w, al), v in zip(kolom, r):
                pdf.cell(w, tinggi, _c(v)[: int(w * 1.9)], border=1, align=al, fill=True)
            pdf.ln()

    bulan = data["bulan"]
    # ── KPI periode terpilih ──
    if periode:
        d = periode["delta"]
        def panah(v, naik_bagus=True, rp=False):
            if v is None:
                return "vs bulan lalu: -"
            s = f"{v:+,.0f}".replace(",", ".") if rp else f"{v:+.1f}".replace(".", ",") + " poin"
            return "vs bulan lalu: " + s
        kpi = [("Nominal belum lunas", rupiah(periode["nominal_belum_lunas"]), panah(d["nominal_belum_lunas"], False, True)),
               ("% Lunas (pelanggan)", _pc(periode["persen_lunas"]), panah(d["persen_lunas"])),
               ("% Pesan dibaca", _pc(periode["persen_baca"]), panah(d["persen_baca"])),
               ("% Pesan dibalas", _pc(periode["persen_balasan"]), panah(d["persen_balasan"]))]
        w = (pdf.w - 20 - 9) / 4
        y = pdf.get_y()
        for i, (k, v, s) in enumerate(kpi):
            x = 10 + i * (w + 3)
            pdf.set_fill_color(*ABU)
            pdf.set_draw_color(203, 213, 225)
            pdf.rect(x, y, w, 20, "DF")
            pdf.set_xy(x + 3, y + 2)
            pdf.set_font("Arial", "", 7)
            pdf.set_text_color(100, 116, 139)
            pdf.cell(w - 6, 4, _c(k))
            pdf.set_xy(x + 3, y + 7)
            pdf.set_font("Arial", "B", 13)
            pdf.set_text_color(*BIRU)
            pdf.cell(w - 6, 6, _c(v))
            pdf.set_xy(x + 3, y + 14.5)
            pdf.set_font("Arial", "", 6.5)
            pdf.set_text_color(100, 116, 139)
            pdf.cell(w - 6, 4, _c(s))
        pdf.set_y(y + 25)
        pdf.set_font("Arial", "", 8)
        pdf.set_text_color(60, 60, 60)
        pdf.cell(0, 5, _c(f"{periode['total']} pelanggan: {periode['lunas']} lunas, {periode['menunggak']} menunggak, "
                         f"{periode['belum_jatuh_tempo']} belum jatuh tempo. Pesan WA terkirim: {periode['wa_terkirim']} "
                         f"({periode['pelanggan_dihubungi']} pelanggan)."), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)

    # ── Grafik batang sederhana (digambar langsung) ──
    def grafik(x, y, w, h, judul_g, nilai, fmt, warna):
        pdf.set_draw_color(203, 213, 225)
        pdf.rect(x, y, w, h)
        pdf.set_xy(x + 3, y + 1.5)
        pdf.set_font("Arial", "B", 8)
        pdf.set_text_color(*BIRU)
        pdf.cell(w - 6, 5, _c(judul_g))
        if not nilai:
            return
        mx = max(v for _, v in nilai) or 1
        ar, ab = w - 10, h - 20
        lb = ar / max(len(nilai), 1)
        for i, (lab, v) in enumerate(nilai):
            bh = ab * (v / mx)
            bx = x + 5 + i * lb + lb * 0.2
            pdf.set_fill_color(*warna)
            pdf.rect(bx, y + h - 9 - bh, lb * 0.6, bh, "F")
            pdf.set_font("Arial", "B", 6)
            pdf.set_text_color(60, 60, 60)
            pdf.set_xy(bx - 3, y + h - 12 - bh)
            pdf.cell(lb * 0.6 + 6, 3, _c(fmt(v)), align="C")
            pdf.set_font("Arial", "", 5.5)
            pdf.set_xy(bx - 4, y + h - 8)
            pdf.cell(lb * 0.6 + 8, 4, _c(lab[:9] + (" " + lab[-2:] if len(lab) > 9 else "")), align="C")

    if bulan:
        y = pdf.get_y()
        gw = (pdf.w - 20 - 4) / 2
        t12 = bulan[-12:]
        pend = lambda l: l.replace("Januari", "Jan").replace("Februari", "Feb").replace("September", "Sep").replace("Agustus", "Agu").replace("Oktober", "Okt").replace("November", "Nov").replace("Desember", "Des")
        grafik(10, y, gw, 52, "% Lunas per bulan", [(pend(b["periode"]), b["persen_lunas"] or 0) for b in t12],
               lambda v: f"{v:.0f}%", (16, 185, 129))
        grafik(10 + gw + 4, y, gw, 52, "Nominal belum lunas per bulan (juta Rp)",
               [(pend(b["periode"]), b["nominal_belum_lunas"] / 1e6) for b in t12], lambda v: f"{v:.1f}", (217, 119, 6))
        pdf.set_y(y + 56)

    # ── Perbandingan antarbulan ──
    judul("Perbandingan antarbulan")
    kolom = [("Periode", 30, "L"), ("Pelanggan", 19, "R"), ("Lunas", 16, "R"), ("Menunggak", 20, "R"), ("Belum JT", 18, "R"),
             ("Nominal belum lunas", 38, "R"), ("% Lunas", 18, "R"), ("Beda % Lunas", 22, "R"), ("WA kirim", 17, "R"),
             ("% Dibaca", 18, "R"), ("% Dibalas", 18, "R"), ("Beda % Baca", 23, "R")]
    baris = []
    for b in reversed(bulan):
        dd = b["delta"]
        baris.append([b["periode"], b["total"], b["lunas"], b["menunggak"], b["belum_jatuh_tempo"],
                      rupiah(b["nominal_belum_lunas"]), _pc(b["persen_lunas"]),
                      "-" if dd["persen_lunas"] is None else f"{dd['persen_lunas']:+.1f}".replace(".", ","),
                      b["wa_terkirim"], _pc(b["persen_baca"]), _pc(b["persen_balasan"]),
                      "-" if dd["persen_baca"] is None else f"{dd['persen_baca']:+.1f}".replace(".", ",")])
    if baris:
        tabel(kolom, baris)
    else:
        pdf.set_font("Arial", "I", 8)
        pdf.cell(0, 6, "Belum ada data.", new_x="LMARGIN", new_y="NEXT")

    # ── Efektivitas template ──
    pdf.ln(4)
    if pdf.get_y() > pdf.h - 60:
        pdf.add_page()
    judul("Efektivitas template WhatsApp")
    if data["template"]:
        kolom_t = [("Template", 66, "L"), ("Tahap", 14, "C"), ("Pesan", 17, "R"), ("Pelanggan", 21, "R"), ("% Dibaca", 20, "R"),
                   ("% Dibalas", 20, "R"), ("Lunas setelahnya", 28, "R"), ("% Konversi", 22, "R"), ("Rata2 hari", 20, "R"),
                   ("Ket.", 49, "L")]
        tabel(kolom_t, [[t["nama"], "-" if t["tahap"] is None else t["tahap"], t["pesan"], t["pelanggan"], _pc(t["persen_baca"]),
                         _pc(t["persen_balasan"]), t["lunas"], _pc(t["konversi"]),
                         "-" if t["rata_hari"] is None else str(t["rata_hari"]).replace(".", ","),
                         ("Terbaik" if t["terbaik"] else "") + (" Data sedikit" if not t["data_cukup"] else "")]
                        for t in data["template"]])
    else:
        pdf.set_font("Arial", "I", 8)
        pdf.cell(0, 6, "Belum ada pesan WhatsApp yang terkirim.", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font("Arial", "I", 6.8)
    pdf.set_text_color(100, 116, 139)
    for c in data["catatan"]:
        pdf.multi_cell(0, 3.6, _c("* " + c), new_x="LMARGIN", new_y="NEXT")

    # ── Daftar pelanggan periode terpilih ──
    if periode and detail:
        pdf.add_page()
        judul(f"Daftar pelanggan - {periode['periode']} ({len(detail)} data, menunggak terlama di atas)")
        kolom_d = [("No", 10, "C"), ("ID Pelanggan", 34, "L"), ("Nama", 62, "L"), ("No. WA", 30, "L"), ("Tagihan", 28, "R"),
                   ("Jatuh tempo", 24, "C"), ("Terlambat", 20, "C"), ("Pembayaran", 34, "C"), ("Status WA", 35, "C")]
        label = {"MENUNGGAK": "Menunggak", "BELUM_JATUH_TEMPO": "Belum JT", "LUNAS": "Lunas"}
        tabel(kolom_d, [[i, d["id_pelanggan"], d.get("nama") or "-", d.get("no_hp") or "-", rupiah(d.get("nominal")),
                         d["jatuh_tempo"].strftime("%d-%m-%Y") if d.get("jatuh_tempo") else "-",
                         "-" if d["hari_terlambat"] is None else f"{d['hari_terlambat']} hr",
                         label[d["kategori"]], d.get("wa_status") or "BELUM"] for i, d in enumerate(detail[:500], 1)])
        if len(detail) > 500:
            pdf.set_font("Arial", "I", 7)
            pdf.cell(0, 6, f"Menampilkan 500 dari {len(detail)} data. Unduh versi Excel untuk daftar lengkap.",
                     new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


# ─────────────────────────────────────────────────────────────
#  Endpoint laporan
# ─────────────────────────────────────────────────────────────
def _siapkan(periode: Optional[str]):
    t, w, info = muat_data()
    data = susun_analitik(t, w, info)
    if not data["bulan"]:
        raise HTTPException(404, "Belum ada data tunggakan. Upload file terlebih dahulu.")
    pilih = _pilih_periode(data["bulan"], periode)
    return data, baris_detail(t, pilih["periode"]), pilih


@router.get("/laporan/excel")
def laporan_excel(periode: Optional[str] = Query(None), admin: dict = Depends(get_current_admin_from_query)):
    data, detail, pilih = _siapkan(periode)
    log_activity(username=admin.get("username", "admin"), role=admin.get("role", "admin"),
                 action="EXPORT_ANALITIK_EXCEL", details=f"Laporan analitik Excel, periode {pilih['periode']}")
    nama = re.sub(r"[^A-Za-z0-9]+", "_", pilih["periode"]).strip("_")
    return Response(content=bangun_excel(data, detail, pilih),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="laporan_tunggakan_{nama}.xlsx"'})


@router.get("/laporan/pdf")
def laporan_pdf(periode: Optional[str] = Query(None), admin: dict = Depends(get_current_admin_from_query)):
    data, detail, pilih = _siapkan(periode)
    log_activity(username=admin.get("username", "admin"), role=admin.get("role", "admin"),
                 action="EXPORT_ANALITIK_PDF", details=f"Laporan analitik PDF, periode {pilih['periode']}")
    nama = re.sub(r"[^A-Za-z0-9]+", "_", pilih["periode"]).strip("_")
    return Response(content=bangun_pdf(data, detail, pilih), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="laporan_tunggakan_{nama}.pdf"'})
